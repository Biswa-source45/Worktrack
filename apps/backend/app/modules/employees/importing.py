import csv
import io
import zipfile
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError
from app.core.security import generate_temp_password, hash_passwords
from app.modules.audit import service as audit
from app.modules.audit.service import AuditCtx
from app.modules.auth.deps import AuthContext
from app.modules.employees.models import STATUS_ACTIVE, Department, Designation, Role, User
from app.modules.employees.schemas import (
    EmployeeCreate,
    ImportCredential,
    ImportResult,
    ImportRowError,
)
from app.modules.employees.service import snapshot

MAX_BYTES = 1_000_000
# Each row costs one Argon2 hash, so the cap keeps an import within a few seconds.
MAX_ROWS = 100
COLUMNS = (
    "emp_code",
    "name",
    "mobile",
    "email",
    "designation",
    "department",
    "role",
    "manager_emp_code",
    "joined_on",
    "field_eligible",
)
REQUIRED = ("emp_code", "name", "mobile", "designation", "role", "joined_on")
_EXAMPLE = "EMP-001,Asha Rao,9876543210,asha@example.com,Engineer,Operations,Field Employee,,"
TEMPLATE_CSV = ",".join(COLUMNS) + "\n" + _EXAMPLE + "2026-01-15,yes\n"


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores 9876543210 as a float
    return str(value).strip()


def _rows_from_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise AppError("INVALID_FILE", "The CSV file must be UTF-8 encoded.", 422) from None
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]


def _rows_from_xlsx(content: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        if sheet is None:
            return []
        return [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError):
        raise AppError("INVALID_FILE", "The file is not a valid Excel workbook.", 422) from None


def parse_file(filename: str, content: bytes) -> list[tuple[int, dict[str, str]]]:
    """Return (spreadsheet row number, values by column name) for each non-empty data row."""
    if len(content) > MAX_BYTES:
        raise AppError("FILE_TOO_LARGE", "The file is larger than 1 MB.", 413)
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension == "csv":
        rows = _rows_from_csv(content)
    elif extension == "xlsx":
        rows = _rows_from_xlsx(content)
    else:
        raise AppError("INVALID_FILE", "Upload a .xlsx or .csv file.", 422)
    if not rows:
        raise AppError("INVALID_FILE", "The file is empty.", 422)
    header = [h.strip().lower().replace(" ", "_") for h in rows[0]]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise AppError("INVALID_FILE", f"Missing column(s): {', '.join(missing)}.", 422)
    data = [
        (number, dict(zip(header, row, strict=False)))
        for number, row in enumerate(rows[1:], start=2)
        if any(row)
    ]
    if len(data) > MAX_ROWS:
        raise AppError("TOO_MANY_ROWS", f"At most {MAX_ROWS} rows per import.", 422)
    return data


def _truthy(value: str) -> bool:
    return value.strip().lower() in {"1", "true", "yes", "y"}


async def import_employees(
    session: AsyncSession,
    actor: AuthContext,
    ctx: AuditCtx,
    filename: str,
    content: bytes,
    dry_run: bool,
) -> ImportResult:
    rows = parse_file(filename, content)

    designations = {
        n.lower(): i
        for i, n in (await session.execute(select(Designation.id, Designation.name))).all()
    }
    departments = {
        n.lower(): i
        for i, n in (await session.execute(select(Department.id, Department.name))).all()
    }
    roles = {r.name.lower(): r for r in (await session.execute(select(Role))).scalars()}
    existing = (
        await session.execute(select(User.id, User.emp_code, User.mobile, User.email, User.status))
    ).all()
    taken_codes = {e.emp_code.lower() for e in existing}
    taken_mobiles = {e.mobile for e in existing}
    taken_emails = {e.email.lower() for e in existing if e.email}
    # Managers: existing active employees or earlier rows of this file (so no cycles).
    managers: dict[str, int | None] = {
        e.emp_code.lower(): e.id for e in existing if e.status == STATUS_ACTIVE
    }

    errors: list[ImportRowError] = []
    valid: list[tuple[EmployeeCreate, str | None]] = []  # (data, manager emp_code)
    for number, raw in rows:
        problems: list[str] = []
        manager_code = raw.get("manager_emp_code", "").strip().lower() or None
        role = roles.get(raw.get("role", "").lower())
        department = raw.get("department", "")
        try:
            data = EmployeeCreate.model_validate(
                {
                    "emp_code": raw.get("emp_code", ""),
                    "name": raw.get("name", ""),
                    "mobile": raw.get("mobile", ""),
                    "email": raw.get("email") or None,
                    "designation_id": designations.get(raw.get("designation", "").lower(), 0),
                    "department_id": departments.get(department.lower()) if department else None,
                    "role_id": role.id if role else 0,
                    "joined_on": raw.get("joined_on", ""),
                    "field_eligible": _truthy(raw.get("field_eligible", "")),
                }
            )
        except ValidationError as exc:
            data = None
            problems.extend(
                f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()
            )

        if data is not None:
            if not data.designation_id:
                problems.append(f"designation: unknown designation '{raw.get('designation', '')}'")
            if department and data.department_id is None:
                problems.append(f"department: unknown department '{department}'")
            if role is None:
                problems.append(f"role: unknown role '{raw.get('role', '')}'")
            elif not set(role.permissions) <= actor.permissions:
                problems.append(f"role: you cannot assign '{role.name}'")
            if manager_code is not None and manager_code not in managers:
                problems.append(f"manager_emp_code: '{raw['manager_emp_code']}' not found")
            if data.emp_code.lower() in taken_codes:
                problems.append("emp_code: already in use")
            if data.mobile in taken_mobiles:
                problems.append("mobile: already in use")
            if data.email and data.email.lower() in taken_emails:
                problems.append("email: already in use")

        if problems:
            errors.extend(ImportRowError(row=number, message=m) for m in problems)
            continue
        assert data is not None  # noqa: S101  # no problems means validation succeeded
        taken_codes.add(data.emp_code.lower())
        taken_mobiles.add(data.mobile)
        if data.email:
            taken_emails.add(data.email.lower())
        managers[data.emp_code.lower()] = None  # known, id assigned on insert
        valid.append((data, manager_code))

    if errors or dry_run:
        return ImportResult(
            dry_run=dry_run,
            total_rows=len(rows),
            created=0,
            errors=errors,
            credentials=[],
        )

    passwords = [generate_temp_password() for _ in valid]
    hashes = await hash_passwords(passwords)
    ids: dict[str, int] = {code: uid for code, uid in managers.items() if uid is not None}
    credentials: list[ImportCredential] = []
    for (data, manager_code), password, password_hash in zip(valid, passwords, hashes, strict=True):
        user = User(
            **data.model_dump(exclude={"password"}),
            password_hash=password_hash,
            must_change_password=True,
        )
        user.manager_id = ids[manager_code] if manager_code else None
        session.add(user)
        await session.flush()
        ids[data.emp_code.lower()] = user.id
        audit.record(session, ctx, "employee.create", "user", user.id, after=snapshot(user))
        credentials.append(
            ImportCredential(emp_code=user.emp_code, name=user.name, temporary_password=password)
        )
    audit.record(
        session, ctx, "employee.import", "user", after={"file": filename, "created": len(valid)}
    )
    await session.commit()
    return ImportResult(
        dry_run=False, total_rows=len(rows), created=len(valid), errors=[], credentials=credentials
    )
